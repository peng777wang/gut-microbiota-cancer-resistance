"""Build the topic trend table (Table 4 replacement) from the Figure 9 run."""

import csv
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = r"C:/72.wxjlx-1/Figure9_output"
XLSX = OUT + "/Table4_topic_trends.xlsx"

# provisional labels, to be confirmed by the two independent raters named in the Methods
PROVISIONAL = {
    0: "Gut microbiota and cancer immunotherapy",
    1: "Colorectal cancer, Fusobacterium nucleatum and treatment resistance",
    2: "Immune checkpoint inhibitors and drug resistance",
    3: "Pancreatic ductal adenocarcinoma and gemcitabine resistance",
    4: "Bacterial infection, phage therapy and antimicrobial resistance",
    5: "Faecal microbiota transplantation for Clostridioides difficile infection",
    6: "Bile acid metabolism and hepatic drug transporters",
    7: "Inflammatory bowel disease and intestinal inflammation",
    8: "Ulcerative colitis and pouchitis",
    9: "Hepatocellular carcinoma and sorafenib resistance",
    10: "Graft-versus-host disease after haematopoietic stem cell transplantation",
}

trends = pd.read_csv(OUT + "/topic_trends.csv", encoding="utf-8-sig")
words = pd.read_csv(OUT + "/topic_topwords.csv", encoding="utf-8-sig")
annual = pd.read_csv(OUT + "/topics_over_time_annual_matrix.csv", index_col=0)

trends = trends.merge(words[["Topic", "Top10"]], on="Topic", how="left")
trends["Label (provisional)"] = trends["Topic"].map(PROVISIONAL)
annual.columns = [str(c) for c in annual.columns]   # topic ids arrive as strings


def series_for(topic):
    column = str(topic)
    return annual[column] if column in annual.columns else None


trends["Documents"] = trends["Topic"].map(
    lambda t: int(series_for(t).sum()) if series_for(t) is not None else None)
trends["First year"] = trends["Topic"].map(
    lambda t: int(series_for(t)[series_for(t) > 0].index[0])
    if series_for(t) is not None else None)
trends["Peak year"] = trends["Topic"].map(
    lambda t: int(series_for(t).idxmax()) if series_for(t) is not None else None)
trends["Peak count"] = trends["Topic"].map(
    lambda t: int(series_for(t).max()) if series_for(t) is not None else None)

trends = trends[["Topic", "Label (provisional)", "Documents", "R2", "P", "Slope",
                 "Trend", "First year", "Peak year", "Peak count", "Top10"]]
trends = trends.sort_values("Slope", ascending=False).reset_index(drop=True)
trends.insert(0, "Rank", range(1, len(trends) + 1))

trends.to_csv(OUT + "/Table4_topic_trends.csv", index=False, encoding="utf-8-sig")

annual_out = annual.copy()
annual_out.insert(0, "Year", annual_out.index)
annual_out.to_csv(OUT + "/Table4_annual_topic_counts.csv", index=False,
                  encoding="utf-8-sig")

summary = pd.DataFrame([{
    "Documents in corpus": 1532,
    "Valid topics": len(trends),
    "Documents in the outlier topic": 341,
    "Outlier share (%)": 22.3,
    "Topics with a significant increase": int((trends["Trend"] == "significant increase").sum()),
    "Topics with no significant trend": int((trends["Trend"] == "no significant trend").sum()),
    "Topics excluded (fewer than five time points)":
        int(trends["Trend"].str.contains("excluded").sum()),
}])

with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
    trends.to_excel(writer, sheet_name="Topic trends", index=False)
    annual_out.to_excel(writer, sheet_name="Annual counts", index=False)
    summary.to_excel(writer, sheet_name="Run summary", index=False)

    book = writer.book
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for sheet in book.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
        for column in sheet.columns:
            width = max(len(str(c.value)) if c.value else 0 for c in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                max(width + 2, 10), 60)
        sheet.freeze_panes = "A2"

print(trends[["Rank", "Topic", "Label (provisional)", "Documents", "R2", "P",
              "Slope", "Trend"]].to_string(index=False))
print()
print("written:", XLSX)
